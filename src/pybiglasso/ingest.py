"""Blockwise conversion from source paths to ordinary Arrow Tables and files."""

from __future__ import annotations

import csv
import fnmatch
import hashlib
import json
import math
import os
import re
import struct
import tempfile
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import TYPE_CHECKING
from xml.parsers import expat

import numpy as np

from .arrow import open_mapped

if TYPE_CHECKING:
    import pyarrow as pa

# Inference candidates deliberately narrower than the fitting input contract.
_TYPES = ("bool", "int8", "uint8", "int16", "float32", "float64")
_INTEGER = re.compile(r"[+-]?\d+\Z")


@dataclass(frozen=True)
class IngestResult:
    """Persistent output and its read-only design/optional response mapping.

    Files are not deleted when this object is collected. This lets a retained
    design or response outlive the result, including on Windows.
    """

    design: pa.Table
    response: np.ndarray | None
    feature_names: tuple[str, ...]
    path: Path


def _arrow():
    try:
        import pyarrow as pa
    except ImportError as error:
        raise ImportError("ingest needs pip install 'pybiglasso[arrow]'") from error
    return pa


def _names(values):
    names = tuple(values)
    if not names or any(not isinstance(name, str) or not name for name in names):
        raise ValueError("Source columns must have nonempty string names.")
    if len(set(names)) != len(names):
        raise ValueError("Source column names must be unique.")
    return names


def _selection(names, response, features, exclude, match):
    if response is not None and response not in names:
        raise ValueError(f"Unknown response column: {response!r}")
    if isinstance(features, str) or isinstance(exclude, str):
        raise ValueError("features and exclude must be sequences of column names.")
    selected = list(names if features is None else features)
    excluded = set(() if exclude is None else exclude)
    if len(set(selected)) != len(selected):
        raise ValueError("Feature names must be unique.")
    unknown = (set(selected) | excluded) - set(names)
    if unknown:
        raise ValueError(f"Unknown column: {sorted(unknown)!r}")
    if match is not None and not isinstance(match, (str, re.Pattern)):
        raise ValueError("match must be a glob string or compiled regular expression.")
    selected = [
        name
        for name in selected
        if name != response
        and name not in excluded
        and (
            match is None
            or (
                bool(match.search(name))
                if isinstance(match, re.Pattern)
                else fnmatch.fnmatchcase(name, match)
            )
        )
    ]
    if not selected:
        raise ValueError("Select at least one feature column.")
    return tuple(selected)


def _numeric(value, column, row):
    """Preserve integer precision before testing any floating representation."""
    if isinstance(value, Decimal):
        if not value.is_finite() or value.adjusted() > 308:
            raise ValueError(f"Column {column!r}, row {row}: nonfinite numeric value.")
        if value == value.to_integral_value():
            value = int(value)
        else:
            converted = float(value)
            if Decimal.from_float(converted) != value:
                raise ValueError(
                    f"Column {column!r}, row {row}: no supported exact dtype."
                )
            value = converted
    if isinstance(value, str):
        text = value.strip()
        try:
            if text.lower() in ("true", "false"):
                value = text.lower() == "true"
            elif _INTEGER.fullmatch(text):
                value = int(text)
            else:
                number = Decimal(text)
                if not number.is_finite() or number.adjusted() > 308:
                    raise ValueError
                # Integral decimal/exponent spellings must not round via f64.
                value = (
                    -0.0
                    if number.is_zero() and number.is_signed()
                    else (
                        int(number)
                        if number == number.to_integral_value()
                        else float(number)
                    )
                )
        except (ValueError, InvalidOperation, OverflowError):
            raise ValueError(
                f"Column {column!r}, row {row}: missing or nonnumeric value."
            ) from None
    if not isinstance(value, (bool, int, float)):
        raise ValueError(f"Column {column!r}, row {row}: missing or nonnumeric value.")
    try:
        finite = math.isfinite(value)
    except OverflowError:
        finite = False
    if not finite:
        raise ValueError(f"Column {column!r}, row {row}: nonfinite numeric value.")
    return value


def _exact(value, dtype):
    if dtype in ("bool", "int8", "uint8", "int16"):
        if isinstance(value, float) and value == 0 and math.copysign(1, value) < 0:
            return False
        lower, upper = {
            "bool": (0, 1),
            "int8": (-128, 127),
            "uint8": (0, 255),
            "int16": (-32768, 32767),
        }[dtype]
        return lower <= value <= upper and int(value) == value
    try:
        with np.errstate(over="ignore", invalid="ignore"):
            cast = float(np.dtype(dtype).type(value))
        return math.isfinite(cast) and cast == value
    except (OverflowError, ValueError):
        return False


def _headers(source, pa, sheet, excel_shared_strings_limit):
    suffix = source.suffix.lower()
    if suffix == ".csv":
        with source.open(newline="", encoding="utf-8-sig") as stream:
            return _names(next(csv.reader(stream), ()))
    if suffix in (".parquet", ".pq"):
        import pyarrow.parquet as pq

        with pq.ParquetFile(source) as reader:
            return _names(reader.schema_arrow.names)
    if suffix in (".feather", ".arrow", ".ipc"):
        with pa.memory_map(str(source), "r") as mapping:
            return _names(pa.ipc.open_file(mapping).schema.names)
    if suffix == ".xlsx":
        workbook = _workbook(source, excel_shared_strings_limit)
        try:
            worksheet = workbook.worksheets[0] if sheet is None else workbook[sheet]
            return _names(next(worksheet.iter_rows(values_only=True), ()))
        finally:
            workbook.close()
    raise ValueError("Source must be CSV, Parquet, Arrow file (Feather v2), or .xlsx.")


def _check_zip_directory(source):
    # ZipFile allocates the directory before it can expose ZipInfo sizes.
    # Bound the EOCD fields first, choosing the same last signature it uses.
    with source.open("rb") as stream:
        stream.seek(0, 2)
        size = stream.tell()
        stream.seek(max(0, size - (65535 + 22 + 20)))
        tail = stream.read()
    offset = tail.rfind(b"PK\x05\x06")
    if offset < 0 or offset + 22 > len(tail):
        raise ValueError("Invalid XLSX ZIP directory.")
    _, disk, start_disk, disk_entries, entries, size, _, comment = struct.unpack_from(
        "<4s4H2LH", tail, offset
    )
    if offset + 22 + comment != len(tail):
        raise ValueError("Invalid XLSX ZIP directory.")
    if (
        disk
        or start_disk
        or disk_entries != entries
        or entries > 1024
        or size > 1 << 20
        or tail[max(0, offset - 20) : offset - 16] == b"PK\x06\x07"
    ):
        raise ValueError("XLSX ZIP directory exceeds the v1 size/entry limits.")


def _shared_string_parts(archive):
    # The reader selects by content type, not by the usual part filename.
    manifest_limit = 1 << 20
    manifest = archive.getinfo("[Content_Types].xml")
    if manifest.file_size > manifest_limit:
        raise ValueError("XLSX content-types manifest exceeds the 1 MiB limit.")
    with archive.open(manifest) as stream:
        data = stream.read(manifest_limit + 1)
    if len(data) > manifest_limit:
        raise ValueError("XLSX content-types manifest exceeds the 1 MiB limit.")
    parts = set()
    parser = expat.ParserCreate(namespace_separator="}")

    def start(name, attributes):
        if (
            name.rsplit("}", 1)[-1] == "Override"
            and attributes.get("ContentType")
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
        ):
            part = attributes.get("PartName", "")
            if not part.startswith("/"):
                raise ValueError("Invalid XLSX shared-string part name.")
            parts.add(part[1:])

    def reject_doctype(*args):
        raise ValueError("XLSX content-types manifest must not contain a DTD.")

    parser.StartElementHandler = start
    parser.StartDoctypeDeclHandler = reject_doctype
    try:
        parser.Parse(data, True)
    except expat.ExpatError as error:
        raise ValueError("Invalid XLSX content-types manifest.") from error
    return parts


def _check_shared_strings(source, limit):
    # ZipInfo.file_size is the uncompressed size recorded in the directory.
    # Do not open/decompress the table merely to decide whether it is too big.
    _check_zip_directory(source)
    with zipfile.ZipFile(source) as archive:
        members = archive.infolist()
        if len(members) > 1024:
            raise ValueError("XLSX ZIP directory exceeds the v1 size/entry limits.")
        by_name = {member.filename: member for member in members}
        if len(by_name) != len(members):
            raise ValueError("XLSX ZIP contains duplicate part names.")
        parts = {
            name for name in by_name if name.rsplit("/", 1)[-1] == "sharedStrings.xml"
        }
        # Keep the cheap conventional-name rejection before manifest decoding.
        size = sum(by_name[name].file_size for name in parts)
        if size <= limit:
            parts.update(_shared_string_parts(archive))
            try:
                size = sum(by_name[name].file_size for name in parts)
            except KeyError as error:
                raise ValueError("XLSX shared-string part is missing.") from error
    if size > limit:
        raise ValueError(
            "XLSX sharedStrings.xml exceeds excel_shared_strings_limit; "
            "convert the workbook to CSV or Parquet, or raise the ceiling explicitly."
        )


def _workbook(source, excel_shared_strings_limit):
    _check_shared_strings(source, excel_shared_strings_limit)
    try:
        import openpyxl
    except ImportError as error:
        raise ImportError(
            "Excel ingest needs pip install 'pybiglasso[ingest]'"
        ) from error
    return openpyxl.load_workbook(
        source, read_only=True, data_only=True, keep_links=False
    )


def _blocks(
    source, columns, all_names, block_size, sheet, pa, excel_shared_strings_limit
):
    """Yield selected columns; only a source batch or block is decoded at once."""
    suffix = source.suffix.lower()
    if suffix == ".csv":
        import pyarrow.csv as pc

        options = pc.ConvertOptions(
            column_types={name: pa.string() for name in all_names},
            include_columns=list(columns),
            strings_can_be_null=False,
        )
        with pc.open_csv(
            source,
            read_options=pc.ReadOptions(block_size=block_size),
            convert_options=options,
        ) as reader:
            for batch in reader:
                yield dict(
                    zip(
                        columns,
                        (column.to_pylist() for column in batch.columns),
                        strict=True,
                    )
                )
    elif suffix in (".parquet", ".pq"):
        import pyarrow.parquet as pq

        with pq.ParquetFile(source) as reader:
            for batch in reader.iter_batches(
                batch_size=block_size, columns=list(columns), use_threads=False
            ):
                yield {name: batch.column(name).to_pylist() for name in columns}
    elif suffix in (".feather", ".arrow", ".ipc"):
        with pa.memory_map(str(source), "r") as mapping:
            reader = pa.ipc.open_file(mapping)
            for index in range(reader.num_record_batches):
                batch = reader.get_batch(index)
                for start in range(0, batch.num_rows, block_size):
                    yield {
                        name: batch.column(name).slice(start, block_size).to_pylist()
                        for name in columns
                    }
    else:
        workbook = _workbook(source, excel_shared_strings_limit)
        try:
            worksheet = workbook.worksheets[0] if sheet is None else workbook[sheet]
            rows = worksheet.iter_rows(values_only=True)
            next(rows, None)
            indices = [all_names.index(name) for name in columns]
            block = {name: [] for name in columns}
            for row in rows:
                for name, index in zip(columns, indices, strict=True):
                    block[name].append(row[index] if index < len(row) else None)
                if len(block[columns[0]]) == block_size:
                    yield block
                    block = {name: [] for name in columns}
            if block[columns[0]]:
                yield block
        finally:
            workbook.close()


def _fingerprint(source):
    stat = source.stat()
    return dict(path=str(source), size=stat.st_size, mtime_ns=stat.st_mtime_ns)


def _check_source(source, fingerprint):
    if _fingerprint(source) != fingerprint:
        raise ValueError("Source changed during ingest.")


def _infer(blocks, columns, overrides):
    candidates = {
        name: ([overrides[name]] if name in overrides else list(_TYPES))
        for name in columns
    }
    count = 0
    for block in blocks:
        length = len(block[columns[0]])
        for name in columns:
            for index, raw in enumerate(block[name], count + 1):
                value = _numeric(raw, name, index)
                # Counts/IDs beyond the supported integer widths use f64,
                # not f32, even when a particular integer happens to fit f32.
                wide_integer = isinstance(value, int) and not -32768 <= value <= 32767
                candidates[name] = [
                    dtype
                    for dtype in candidates[name]
                    if _exact(value, dtype)
                    and not (
                        name not in overrides and wide_integer and dtype == "float32"
                    )
                ]
                if not candidates[name]:
                    raise ValueError(
                        f"Column {name!r}, row {index}: no supported exact dtype."
                    )
        count += length
    if not count:
        raise ValueError("Source must contain at least one data row.")
    return {name: candidates[name][0] for name in columns}, count


def _stage(blocks, columns, types, directory, count):
    tails = {
        name: np.empty(0, dtype=np.bool_) for name in columns if types[name] == "bool"
    }
    paths = {
        name: directory / f"column-{index}.bin" for index, name in enumerate(columns)
    }
    seen = 0
    for block in blocks:
        for name in columns:
            values = [
                _numeric(raw, name, seen + index + 1)
                for index, raw in enumerate(block[name])
            ]
            if not all(_exact(value, types[name]) for value in values):
                raise ValueError(
                    f"Column {name!r}: source values changed during ingest."
                )
            array = np.asarray(values, dtype=types[name])
            if types[name] == "bool":
                bits = np.concatenate((tails[name], array))
                complete = len(bits) // 8 * 8
                array = np.packbits(bits[:complete], bitorder="little")
                tails[name] = bits[complete:].copy()
            with paths[name].open("ab") as stream:
                stream.write(array.tobytes())
        seen += len(block[columns[0]])
    if seen != count:
        raise ValueError("Source row count changed during ingest.")
    for name, tail in tails.items():
        if len(tail):
            with paths[name].open("ab") as stream:
                stream.write(np.packbits(tail, bitorder="little").tobytes())
    return paths


def _open_result(path, metadata):
    response = None
    if metadata["response"] is not None:
        response = np.memmap(
            path.with_suffix(path.suffix + ".response"),
            mode="r",
            dtype=np.float64,
            shape=(metadata["rows"],),
        )
    design = open_mapped(path)
    if design.shape != (
        metadata["rows"],
        len(metadata["features"]),
    ) or tuple(design.column_names) != tuple(metadata["features"]):
        raise ValueError("Ingest cache metadata does not match its design.")
    return IngestResult(design, response, tuple(design.column_names), path)


def _write_ipc(path, raw, selected, types, count, pa):
    # Keep mmap-backed arrays local to this call: release every staging view
    # before TemporaryDirectory removes its files, including on Windows.
    arrays = []
    arrow_types = {
        "bool": pa.bool_(),
        "int8": pa.int8(),
        "uint8": pa.uint8(),
        "int16": pa.int16(),
        "float32": pa.float32(),
        "float64": pa.float64(),
    }
    for name in selected:
        buffer = pa.memory_map(str(raw[name]), "r").read_buffer()
        arrays.append(
            pa.Array.from_buffers(arrow_types[types[name]], count, [None, buffer])
        )
    batch = pa.record_batch(arrays, names=list(selected))
    with (
        pa.OSFile(str(path), "wb") as stream,
        pa.ipc.new_file(
            stream, batch.schema, options=pa.ipc.IpcWriteOptions(compression=None)
        ) as writer,
    ):
        writer.write_batch(batch)


def ingest(
    source,
    out=None,
    *,
    response=None,
    features=None,
    exclude=None,
    match=None,
    dtypes="auto",
    block_size=65536,
    sheet=None,
    excel_shared_strings_limit=1 << 30,
    cache=False,
    cache_dir=None,
):
    """Convert a numeric source to persistent, feature-only Arrow file storage.

    ``match`` is a glob string or compiled regex. ``dtypes`` is ``'auto'`` or
    a partial mapping of feature names to exact supported physical types.
    Response selection is explicit and produces a separate read-only f64
    mapping. No response/family is inferred from names. Error row numbers are
    one-based data rows (excluding the header).

    ``block_size`` means bytes for CSV and rows for other readers. Compressed
    Feather decoding is bounded by its largest existing record batch. Excel
    is an exception to block-bounded memory: its complete workbook shared-string
    table may be held in memory, including unselected columns. Before each open,
    ``excel_shared_strings_limit`` caps its declared uncompressed XML size
    (1 GiB by default), without decompressing it. This is not an RSS ceiling.

    Specify an explicit persistent ``out`` or ``cache_dir``. There is no
    implicit temporary destination or automatic deletion. ``cache=True``
    reuses matching source path/size/mtime/options; otherwise a fresh artifact
    is written. Explicit existing outputs are never overwritten. Callers own
    the returned output and its .response/.json sidecars.
    """
    if out is None and cache_dir is None:
        raise ValueError("Specify a persistent out or cache_dir for ingest.")
    if not isinstance(source, (str, os.PathLike)):
        raise TypeError(
            "ingest accepts a filesystem path, not a Table, receipt, stream or file object."
        )
    pa = _arrow()
    source = Path(source).resolve(strict=True)
    if (
        not isinstance(block_size, int)
        or isinstance(block_size, bool)
        or block_size <= 0
    ):
        raise ValueError("block_size must be a positive integer.")
    if (
        not isinstance(excel_shared_strings_limit, int)
        or isinstance(excel_shared_strings_limit, bool)
        or excel_shared_strings_limit <= 0
    ):
        raise ValueError("excel_shared_strings_limit must be a positive integer.")
    names = _headers(source, pa, sheet, excel_shared_strings_limit)
    selected = _selection(names, response, features, exclude, match)
    if isinstance(dtypes, str):
        if dtypes != "auto":
            raise ValueError("dtypes must be 'auto' or a mapping.")
        overrides = {}
    elif isinstance(dtypes, Mapping):
        overrides = {name: np.dtype(dtype).name for name, dtype in dtypes.items()}
    else:
        raise ValueError("dtypes must be 'auto' or a mapping.")
    if set(overrides) - set(selected) or any(
        dtype not in _TYPES for dtype in overrides.values()
    ):
        raise ValueError(
            "dtype overrides must name selected features and supported types."
        )
    columns = (*selected, *((response,) if response is not None else ()))
    fingerprint = _fingerprint(source)
    options = dict(
        version=4,
        source=fingerprint,
        features=selected,
        response=response,
        dtypes=overrides,
        sheet=sheet,
        block_size=block_size,
        excel_shared_strings_limit=excel_shared_strings_limit,
    )
    key = hashlib.sha256(json.dumps(options, sort_keys=True).encode()).hexdigest()
    if out is None:
        root = Path(cache_dir).absolute()
        root.mkdir(parents=True, exist_ok=True)
        directory = (
            root / key if cache else Path(tempfile.mkdtemp(prefix="design-", dir=root))
        )
        directory.mkdir(exist_ok=True)
        path = directory / "design.arrow"
    else:
        path = Path(out).absolute()
        path.parent.mkdir(parents=True, exist_ok=True)
    metadata_path = path.with_suffix(path.suffix + ".json")
    response_path = path.with_suffix(path.suffix + ".response")
    if cache and path.exists() and metadata_path.exists():
        metadata = json.loads(metadata_path.read_text())
        if metadata.get("key") == key:
            return _open_result(path, metadata)
    if any(p.exists() for p in (path, metadata_path, response_path)):
        raise FileExistsError(f"Ingest output already exists: {path}")

    def blocks():
        return _blocks(
            source, columns, names, block_size, sheet, pa, excel_shared_strings_limit
        )

    # The response is separate solver input, always exactly representable f64.
    types, count = _infer(
        blocks(),
        columns,
        {**overrides, **({response: "float64"} if response is not None else {})},
    )
    _check_source(source, fingerprint)
    with tempfile.TemporaryDirectory(prefix=".ingest-", dir=path.parent) as temporary:
        staging = Path(temporary)
        raw = _stage(blocks(), columns, types, staging, count)
        _check_source(source, fingerprint)
        output = staging / "design.arrow"
        _write_ipc(output, raw, selected, types, count, pa)
        metadata = dict(key=key, rows=count, features=selected, response=response)
        staged_metadata = staging / "metadata.json"
        staged_metadata.write_text(json.dumps(metadata, sort_keys=True))
        # Hard links publish without overwriting an existing destination; all
        # staging files are on this filesystem. Roll back only our own links.
        published = []
        try:
            for origin, destination in [
                *(([(raw[response], response_path)]) if response is not None else []),
                (output, path),
                (staged_metadata, metadata_path),
            ]:
                os.link(origin, destination)
                published.append(destination)
        except OSError:
            for destination in published:
                destination.unlink()
            raise
    return _open_result(path, metadata)
